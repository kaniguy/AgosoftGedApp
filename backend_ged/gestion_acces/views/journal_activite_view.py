from io import BytesIO

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated

from gestion_acces.models.journal_activite import JournalActivite
from gestion_acces.serializers.journal_activite_serializer import JournalActiviteSerializer


class JournalActivitePagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 200


class JournalActiviteViewSet(viewsets.ReadOnlyModelViewSet):
    """Consultation du journal d'activité (lecture seule)."""

    serializer_class = JournalActiviteSerializer
    permission_classes = [IsAuthenticated]
    pagination_class = JournalActivitePagination

    def get_queryset(self):
        qs = JournalActivite.objects.select_related("utilisateur").all()
        user = self.request.user

        if not (
            user.is_superuser
            or user.has_perm("gestion_acces.view_journalactivite")
        ):
            return JournalActivite.objects.none()

        params = self.request.query_params

        action = (params.get("action") or "").strip()
        if action:
            qs = qs.filter(action=action)

        categorie = (params.get("categorie") or "").strip()
        if categorie:
            qs = qs.filter(categorie=categorie)

        search = (params.get("search") or "").strip()
        if search:
            qs = qs.filter(
                Q(nom_utilisateur__icontains=search)
                | Q(description__icontains=search)
                | Q(chemin__icontains=search)
                | Q(adresse_ip__icontains=search)
                | Q(utilisateur__username__icontains=search)
                | Q(utilisateur__first_name__icontains=search)
                | Q(utilisateur__last_name__icontains=search)
                | Q(utilisateur__email__icontains=search)
            )

        date_debut = (params.get("date_debut") or "").strip()
        if date_debut:
            qs = qs.filter(date_creation__date__gte=date_debut)

        date_fin = (params.get("date_fin") or "").strip()
        if date_fin:
            qs = qs.filter(date_creation__date__lte=date_fin)

        return qs

    @action(detail=False, methods=["get"], url_path="export-excel")
    def export_excel(self, request):
        """Exporte le journal filtré au format Excel (.xlsx)."""
        qs = self.get_queryset().order_by("-date_creation")[:5000]

        wb = Workbook()
        ws = wb.active
        ws.title = "Journal d'activité"

        headers = [
            "Date / heure",
            "Prénom",
            "Nom",
            "Nom d'utilisateur",
            "Email",
            "Action",
            "Catégorie",
            "Description",
            "IP",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")

        for row in qs:
            user = row.utilisateur
            prenom = (user.first_name if user else "") or ""
            nom = (user.last_name if user else "") or ""
            email = (user.email if user else "") or ""
            local_dt = timezone.localtime(row.date_creation) if row.date_creation else None
            date_str = local_dt.strftime("%d/%m/%Y %H:%M:%S") if local_dt else ""
            ws.append(
                [
                    date_str,
                    prenom,
                    nom,
                    row.nom_utilisateur or (user.username if user else ""),
                    email,
                    row.get_action_display(),
                    row.get_categorie_display(),
                    row.description or "",
                    row.adresse_ip or "",
                ]
            )

        widths = [20, 14, 14, 18, 28, 16, 18, 50, 16]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"journal-activite-{timezone.localtime().strftime('%Y%m%d-%H%M%S')}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
